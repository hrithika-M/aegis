# -*- coding: utf-8 -*-
"""POD2 · Step 1 — isolate the clean June load-factor data.

The product owner (Avra) shared 'JUN LOADS FOR AIRLINE.xlsx' with two sheets. Per
the walkthrough call, only the JUNE data is trusted real actuals; the second sheet
(labelled July) is calculated. We therefore use ONLY the 'June Acft' sheet — real,
hardcoded daily load factors by aircraft category:

  Category C = narrow-body (A320/A321/B737, operated by 6E/AI/IX/TR)
  Category B = ATR-72 (6E)
  Category A = 9-seater (India One / I7)

Output: june_load_daily.csv  (long/tidy format, one row per date × category × direction)
Run:    python extract_june.py
"""
import csv, os
import datetime as dt
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, 'inputs', 'JUN_LOADS_FOR_AIRLINE.xlsx')
OUT = os.path.join(HERE, 'june_load_daily.csv')

# category -> (label, arr_lf_col, arr_pax_col, dep_lf_col, dep_pax_col)   [1-indexed]
CATS = {
    'C': ('Narrow-body (A320/A321/B737)', 2, 3, 4, 5),
    'B': ('ATR-72', 6, 7, 8, 9),
    'A': ('9-seater (India One)', 10, 11, 12, 13),
}
DOW = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    sheet = next(s for s in wb.sheetnames if s.strip().lower() == 'load factor sheet june acft')
    ws = wb[sheet]
    rows = []
    for r in range(4, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if not isinstance(d, dt.datetime):
            break                                  # stop at the TOTAL row
        for cat, (label, alf, apx, dlf, dpx) in CATS.items():
            for direction, lf_c, px_c in [('ARR', alf, apx), ('DEP', dlf, dpx)]:
                lf = ws.cell(r, lf_c).value
                px = ws.cell(r, px_c).value
                rows.append({
                    'date': d.date().isoformat(), 'day_of_week': DOW[d.weekday()],
                    'category': cat, 'aircraft_group': label, 'direction': direction,
                    'load_factor': round(float(lf), 4) if isinstance(lf, (int, float)) else 0,
                    'pax': int(px) if isinstance(px, (int, float)) else 0,
                })
    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=['date', 'day_of_week', 'category', 'aircraft_group',
                                          'direction', 'load_factor', 'pax'])
        w.writeheader(); w.writerows(rows)

    days = sorted({r['date'] for r in rows})
    dep = sum(r['pax'] for r in rows if r['direction'] == 'DEP')
    arr = sum(r['pax'] for r in rows if r['direction'] == 'ARR')
    print(f"june_load_daily.csv — {len(rows)} rows, {len(days)} days ({days[0]} .. {days[-1]})")
    print(f"  total June passengers: {dep:,} departing + {arr:,} arriving = {dep + arr:,}")
    for cat, (label, *_ ) in CATS.items():
        cp = sum(r['pax'] for r in rows if r['category'] == cat)
        print(f"  category {cat} {label:<32} {cp:>7,} pax")


if __name__ == '__main__':
    main()
