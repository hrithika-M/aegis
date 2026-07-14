# -*- coding: utf-8 -*-
"""The end-to-end calibrated hourly demand curve for VTZ.

Folds every real input we gathered into one typical-week hourly grid:

  real monthly passengers PER ROUTE, PER DIRECTION (DGCA, avg 2025)
     distributed across that route's actual flight instances in the schedule
     (which flights, what hour, what days)  ->  hourly arrivals & departures
  => each route's weekly total x 4.345 == its real monthly DGCA number, by construction.

No flat load-factor guess: passengers come straight from DGCA route counts, so the
implied load factor is each route's real one. Routes with no DGCA data (international
AUH/SIN, or the bad code 881) fall back to seats x 0.82, clearly the only assumed part.

Run:  python build_hourly.py  -> ../data/bhogapuram_vtz/vtz_hourly.csv
"""
import csv, os
from collections import defaultdict
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, '..', 'data', 'bhogapuram_vtz')
SCHED = os.path.join(DATA, 'VTZ_schedule_S26.xlsx')
CITYPAIR = os.path.join(DATA, 'vtz_citypair_monthly.csv')
OUT = os.path.join(DATA, 'vtz_hourly.csv')
WK = 4.345                                   # avg weeks per month
DOW = ['Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun']
CODE2CITY = {'HYD': 'Hyderabad', 'DEL': 'Delhi', 'BLR': 'Bengaluru', 'MAA': 'Chennai',
             'BOM': 'Mumbai', 'NMI': 'Mumbai', 'CCU': 'Kolkata', 'VGA': 'Vijayawada',
             'TIR': 'Tirupati', 'KJB': 'Kurnool', 'PYB': 'Jeypore', 'BBI': 'Bhubaneswar'}


def freq_days(fr):
    return {int(c) - 1 for c in str(fr) if c.isdigit() and c != '0'}


def hour_of(t):
    return int(str(t).split(':')[0])


def real_route_monthly():
    """avg 2025 monthly pax per route, directional (arr to VTZ / dep from VTZ)."""
    arr, dep, n = defaultdict(float), defaultdict(float), defaultdict(set)
    for r in csv.DictReader(open(CITYPAIR, encoding='utf-8')):
        if int(r['Year']) != 2025:
            continue
        c = r['Route']
        arr[c] += float(r['PaxArr_toVTZ'] or 0)
        dep[c] += float(r['PaxDep_fromVTZ'] or 0)
        n[c].add((r['Year'], r['Month']))
    arr = {c: v / len(n[c]) for c, v in arr.items()}
    dep = {c: v / len(n[c]) for c, v in dep.items()}
    return arr, dep


def main():
    arr_pax, dep_pax = real_route_monthly()
    ws = openpyxl.load_workbook(SCHED, data_only=True).active

    # collect flight instances; first count weekly instances per route/direction
    arr_fl, dep_fl = [], []            # (city, hour, freqset, seats)
    wk_arr, wk_dep = defaultdict(int), defaultdict(int)
    for r in range(3, ws.max_row + 1):
        sl = ws.cell(r, 1).value
        if not isinstance(sl, int):
            break
        seats = ws.cell(r, 11).value
        if not seats:
            continue
        seats = int(seats); fr = freq_days(ws.cell(r, 2).value)
        oc = CODE2CITY.get(str(ws.cell(r, 5).value).strip().upper())          # arrival origin
        dc = CODE2CITY.get(str(ws.cell(r, 8).value).strip().upper())          # departure dest
        arr_fl.append((oc, hour_of(ws.cell(r, 6).value), fr, seats, str(ws.cell(r, 5).value).strip()))
        dep_fl.append((dc, hour_of(ws.cell(r, 9).value), fr, seats, str(ws.cell(r, 8).value).strip()))
        if oc:
            wk_arr[oc] += len(fr)
        if dc:
            wk_dep[dc] += len(fr)

    def per_instance(city, seats, wk_count, real):
        if city and city in real and wk_count:
            return real[city] / (wk_count * WK)          # real DGCA pax / monthly instances
        return seats * 0.82                              # fallback (intl / bad code)

    grid_dep = defaultdict(float)
    grid_arr = defaultdict(float)
    for city, h, fr, seats, _ in arr_fl:
        p = per_instance(city, seats, wk_arr.get(city, 0), arr_pax)
        for d in fr:
            grid_arr[(d, h)] += p
    for city, h, fr, seats, _ in dep_fl:
        p = per_instance(city, seats, wk_dep.get(city, 0), dep_pax)
        for d in fr:
            grid_dep[(d, h)] += p

    rows = []
    for d in range(7):
        for h in range(24):
            dep, arr = round(grid_dep[(d, h)]), round(grid_arr[(d, h)])
            if dep or arr:
                rows.append([DOW[d], f"{h:02d}:00", dep, arr, dep + arr])
    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['DayOfWeek', 'Hour', 'DepPAX', 'ArrPAX', 'TotalPAX'])
        w.writerows(rows)

    wk_total = sum(r[4] for r in rows)
    print(f"wrote {os.path.basename(OUT)}  ({len(rows)} populated hour-slots)")
    print(f"implied MONTHLY total: {round(wk_total * WK):,}  (real 2025 avg ~240,000)")
    # peak hours on a typical Wednesday
    wed = [r for r in rows if r[0] == 'Wed']
    pk_dep = max(wed, key=lambda r: r[2]); pk_arr = max(wed, key=lambda r: r[3])
    print(f"Wed peak departures {pk_dep[1]} = {pk_dep[2]} boarding")
    print(f"Wed peak arrivals   {pk_arr[1]} = {pk_arr[3]} deplaning")


if __name__ == '__main__':
    main()
