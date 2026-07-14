# -*- coding: utf-8 -*-
"""Build (1) a real national daily reference and (2) a VTZ daily PAX estimate.

VTZ-specific daily counts are not published anywhere. So we DERIVE them, anchored
to real data:

  VTZ daily pax = real VTZ monthly total (DGCA)  distributed across the month by
     weight(day) = schedule seats operating that weekday   (VTZ-specific rhythm)
                 x national daily index for that date       (real calendar effects:
                                                              holidays, disruptions)
  scaled so each month sums EXACTLY to the real DGCA monthly total.

Inputs (committed):
  ../data/bhogapuram_vtz/VTZ_schedule_S26.xlsx     -> weekday seat weights
  ../data/bhogapuram_vtz/vtz_citypair_monthly.csv  -> real VTZ monthly totals
  <scratch>/daily.csv (MoCA national daily)         -> national daily shape + LFs

Outputs:
  ../data/bhogapuram_vtz/national_daily_reference.csv   (REAL, all-India)
  ../data/bhogapuram_vtz/vtz_daily_estimate.csv         (DERIVED, VTZ)

Run:  python daily_estimate.py
"""
import csv, datetime as dt, os
from collections import defaultdict
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, '..', 'data', 'bhogapuram_vtz')
SCHED = os.path.join(DATA, 'VTZ_schedule_S26.xlsx')
CITYPAIR = os.path.join(DATA, 'vtz_citypair_monthly.csv')
# national daily is fetched to scratch (too big/national to commit whole); trimmed below
DAILY_SRC = os.path.join(HERE, 'daily.csv')
NAT_OUT = os.path.join(DATA, 'national_daily_reference.csv')
VTZ_OUT = os.path.join(DATA, 'vtz_daily_estimate.csv')
TARGET_YEAR = 2025                       # full year we have real VTZ monthly totals for
DOW = ['Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun']


def freq_days(fr):
    return {int(c) - 1 for c in str(fr) if c.isdigit() and c != '0'}


def weekday_seat_weights():
    """From the schedule: total seats (arr+dep footfall) operating each weekday."""
    ws = openpyxl.load_workbook(SCHED, data_only=True).active
    w = {d: 0 for d in range(7)}
    for r in range(3, ws.max_row + 1):
        sl = ws.cell(r, 1).value
        if not isinstance(sl, int):
            break
        seats = ws.cell(r, 11).value
        if not seats:
            continue
        for d in freq_days(ws.cell(r, 2).value):
            w[d] += 2 * int(seats)           # arrival + departure footfall
    return w


def build_national_reference():
    # Reproducible fallback: if the raw MoCA daily.csv isn't present, read the
    # already-committed trimmed reference instead of re-fetching.
    if not os.path.exists(DAILY_SRC):
        natpax = {}
        for r in csv.DictReader(open(NAT_OUT, encoding='utf-8')):
            natpax[r['Date']] = int(r['DomesticPax'])
        return natpax
    rows = list(csv.DictReader(open(DAILY_SRC, encoding='utf-8')))
    out = []
    for x in rows:
        try:
            dom = int(x['Domestic (Arriving Pax)'] or 0) + int(x['Domestic (Departing Pax)'] or 0)
            intl = int(x['International (Arriving Pax)'] or 0) + int(x['International (Departing Pax)'] or 0)
        except ValueError:
            continue
        if dom <= 0:
            continue
        d = dt.date.fromisoformat(x['Date'])
        out.append([x['Date'], DOW[d.weekday()], dom, intl,
                    x['Passenger Load Factor (Indigo)'].strip('%') or '',
                    x['Passenger Load Factor (Air India)'].strip('%') or '',
                    x['Passenger Load Factor (Akasa Air)'].strip('%') or '',
                    x['Passenger Load Factor (Spicejet)'].strip('%') or ''])
    with open(NAT_OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['Date', 'DayOfWeek', 'DomesticPax', 'InternationalPax',
                    'LF_IndiGo_pct', 'LF_AirIndia_pct', 'LF_Akasa_pct', 'LF_SpiceJet_pct'])
        w.writerows(out)
    # return date -> domestic pax for the index
    return {r[0]: r[2] for r in out}


def real_vtz_monthly():
    tot = defaultdict(float)
    for r in csv.DictReader(open(CITYPAIR, encoding='utf-8')):
        tot[(int(r['Year']), int(r['Month']))] += float(r['PaxTotal'] or 0)
    return tot


def main():
    wsw = weekday_seat_weights()
    natpax = build_national_reference()
    monthly = real_vtz_monthly()

    # national monthly means (for the daily index)
    nat_month = defaultdict(list)
    for ds, p in natpax.items():
        d = dt.date.fromisoformat(ds)
        nat_month[(d.year, d.month)].append(p)
    nat_mean = {k: (sum(v) / len(v)) for k, v in nat_month.items()}

    # raw weights per day of TARGET_YEAR
    rows = []
    d = dt.date(TARGET_YEAR, 1, 1)
    while d.year == TARGET_YEAR:
        idx = natpax.get(d.isoformat())
        nat_idx = (idx / nat_mean[(d.year, d.month)]) if idx and nat_mean.get((d.year, d.month)) else 1.0
        raw = wsw[d.weekday()] * nat_idx
        rows.append([d, raw])
        d += dt.timedelta(days=1)

    # scale each month to the real DGCA total
    month_raw = defaultdict(float)
    for dd, raw in rows:
        month_raw[(dd.year, dd.month)] += raw
    out = []
    for dd, raw in rows:
        key = (dd.year, dd.month)
        real_tot = monthly.get(key, 0)
        est = raw / month_raw[key] * real_tot if month_raw[key] else 0
        out.append([dd.isoformat(), DOW[dd.weekday()], round(est), int(real_tot)])
    with open(VTZ_OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['Date', 'DayOfWeek', 'VTZ_EstPax', 'RealMonthlyTotal_DGCA'])
        w.writerows(out)

    print(f"wrote {os.path.basename(NAT_OUT)}  ({len(natpax)} real national days)")
    print(f"wrote {os.path.basename(VTZ_OUT)}  ({len(out)} VTZ daily estimates, {TARGET_YEAR})")
    # sanity: a month should sum to its real total
    jan = [r for r in out if r[0].startswith(f'{TARGET_YEAR}-01')]
    print(f"  check Jan {TARGET_YEAR}: sum of daily = {sum(r[2] for r in jan):,} vs real {jan[0][3]:,}")
    print(f"  sample: busiest & quietest day in Jan:")
    js = sorted(jan, key=lambda r: r[2])
    print(f"    {js[-1][0]} ({js[-1][1]}) {js[-1][2]:,}   |   {js[0][0]} ({js[0][1]}) {js[0][2]:,}")


if __name__ == '__main__':
    main()
