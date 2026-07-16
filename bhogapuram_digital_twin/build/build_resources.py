# -*- coding: utf-8 -*-
"""Resource master — the REAL airport numbering from MASTER DATA RESOURCES.xlsx.

Meeting feedback (Bhogapuram, Jul-2026): "random figures connected to the airport —
gates numbers and other numbering — change them to realistic figures, numbers and
gates." This replaces generic counts with the airport's actual inventory:

  Stands  (11L / 11 / 11R / 12L ... with apron, ICAO code, contact/remote)
  Gates   (A2R / A2 / A2L / A4 / A6 ... with traffic type)
  Belts   (1..4 with traffic type)
  CIC     check-in positions expanded from ranges (A01-A10, B01-B10, ...) with type
          (CONVENTIONAL counter vs SBD self-bag-drop)

Output: master/Resource_Master.csv  (one row per physical resource)
Run:    python build_resources.py
"""
import csv, os, re
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, '..', 'data', 'bhogapuram_vtz', 'master_data_resources.xlsx')
OUT = os.path.join(ROOT, 'master', 'Resource_Master.csv')


def rows_of(ws):
    """Yield non-empty data rows (cells B..H), skipping the two header rows."""
    for r in range(4, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(2, 9)]
        if any(v is not None and str(v).strip() != '' for v in row):
            yield [str(v).strip() if v is not None else '' for v in row]


def expand_range(rng):
    """'A01-A10' -> A01..A10 ; tolerates 'B09-B10' etc."""
    m = re.match(r'([A-Z])\s*0*(\d+)\s*-\s*[A-Z]?\s*0*(\d+)', str(rng).replace(' ', ''))
    if not m:
        return [str(rng).strip()]
    a, lo, hi = m.groups()
    return [f"{a}{i:02d}" for i in range(int(lo), int(hi) + 1)]


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    sheets = {n.strip().lower(): wb[n] for n in wb.sheetnames}
    out = []
    for row in rows_of(sheets['stands']):     # SL, STAND, APRON, Traffic, ICAO, CONTACT/REMOTE
        out.append(['STAND', row[1], row[2], row[3], row[4], row[5]])
    for row in rows_of(sheets['gates']):      # SL, GATE, APRON, Traffic, CONTACT/REMOTE
        out.append(['GATE', row[1], row[2], row[3], '', row[4]])
    for row in rows_of(sheets['belts']):      # SL, BELT, Traffic
        out.append(['BELT', f"Belt {row[1]}", '', row[2], '', ''])
    for row in rows_of(sheets['cic']):        # SL, AREA, RANGE, Type
        for pos in expand_range(row[2]):
            out.append(['CHECKIN', pos, f"Area {row[1].strip()}", '', row[3], ''])

    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['resource_type', 'resource_id', 'zone', 'traffic_type', 'spec', 'contact_remote'])
        w.writerows(out)

    counts = {}
    for r in out:
        counts[r[0]] = counts.get(r[0], 0) + 1
    print(f"Resource_Master.csv — real airport inventory ({len(out)} resources)")
    for k, v in sorted(counts.items()):
        print(f"  {k:<8} {v}")
    sbd = sum(1 for r in out if r[0] == 'CHECKIN' and 'SBD' in (r[4] or '').upper())
    print(f"  (check-in split: {sbd} SBD self-bag-drop / {counts.get('CHECKIN', 0) - sbd} conventional)")


if __name__ == '__main__':
    main()
