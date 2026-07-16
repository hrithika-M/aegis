# -*- coding: utf-8 -*-
"""Pattern store — ingest every data drop, learn its patterns, NEVER restart from zero.

Meeting requirement (Bhogapuram, Jul-2026): "data is shared min of 20 days before —
extract the records that we get and then reuse it, don't come to zero again;
patterns should be found and reusing of those patterns should be done."

Every file the airport shares (schedule versions, later real ops extracts) is
ingested ONCE (fingerprinted by content hash), its patterns are extracted, and the
store accumulates — so each new drop refines the picture instead of resetting it.

Patterns extracted from a schedule drop:
  - weekly movement pattern (flights per day-of-week)
  - hourly bank structure (dep/arr movements per hour)
  - route mix (weekly seats per destination)
  - fleet mix (movements per aircraft type)

Store: patterns/patterns.json  (snapshots keyed by content hash + a merged
'current' view that always reflects the latest drop per source type, with the
full history retained for drift comparison).

Usage:
  python pattern_store.py ingest <path-to-drop.xlsx>
  python pattern_store.py show
Seeded with the S26 schedule on first run (no args).
"""
import hashlib, json, os, sys
import datetime as dt
from collections import defaultdict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
STORE = os.path.join(HERE, 'patterns.json')
DEFAULT_DROP = os.path.join(ROOT, '..', 'data', 'bhogapuram_vtz', 'VTZ_schedule_S26.xlsx')


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(65536), b''):
            h.update(chunk)
    return h.hexdigest()[:16]


def load_store():
    if os.path.exists(STORE):
        return json.load(open(STORE, encoding='utf-8'))
    return {'snapshots': [], 'current': {}}


def save_store(s):
    with open(STORE, 'w', encoding='utf-8') as f:
        json.dump(s, f, indent=2)


def extract_schedule_patterns(path):
    """Patterns from a VTZ-style schedule workbook ('MAIN FILE' sheet)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    name = 'MAIN FILE' if 'MAIN FILE' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[name]
    dow_moves = defaultdict(int)
    dep_hour = defaultdict(int); arr_hour = defaultdict(int)
    route_seats = defaultdict(int); fleet = defaultdict(int)
    n = 0
    for r in range(3, ws.max_row + 1):
        if not isinstance(ws.cell(r, 1).value, int):
            break
        seats = ws.cell(r, 11).value or 0
        freq = [int(c) - 1 for c in str(ws.cell(r, 2).value) if c.isdigit() and c != '0']
        try:
            ah = int(str(ws.cell(r, 6).value).split(':')[0])
            dh = int(str(ws.cell(r, 9).value).split(':')[0])
        except (ValueError, AttributeError):
            continue
        n += 1
        dest = str(ws.cell(r, 8).value).strip()
        acft = str(ws.cell(r, 3).value).strip()
        for d in freq:
            dow_moves[d] += 2                       # arr + dep
            route_seats[dest] += int(seats)
            fleet[acft] += 2
        dep_hour[dh] += len(freq)
        arr_hour[ah] += len(freq)
    dows = ['Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun']
    return {
        'source_type': 'schedule',
        'rotations': n,
        'weekly_movements_by_dow': {dows[k]: v for k, v in sorted(dow_moves.items())},
        'weekly_dep_by_hour': {f"{k:02d}": v for k, v in sorted(dep_hour.items())},
        'weekly_arr_by_hour': {f"{k:02d}": v for k, v in sorted(arr_hour.items())},
        'weekly_seats_by_destination': dict(sorted(route_seats.items(), key=lambda x: -x[1])),
        'weekly_movements_by_aircraft': dict(sorted(fleet.items(), key=lambda x: -x[1])),
    }


def ingest(path):
    store = load_store()
    fp = sha(path)
    if any(s['sha16'] == fp for s in store['snapshots']):
        print(f"already ingested ({fp}) — store unchanged ({len(store['snapshots'])} snapshots)")
        return
    pat = extract_schedule_patterns(path)
    snap = {'sha16': fp, 'file': os.path.basename(path),
            'ingested_at': dt.datetime.now().replace(microsecond=0).isoformat(),
            'patterns': pat}
    store['snapshots'].append(snap)
    store['current'][pat['source_type']] = {'sha16': fp, 'file': snap['file'],
                                            'ingested_at': snap['ingested_at'],
                                            'patterns': pat}
    save_store(store)
    print(f"ingested {snap['file']} ({fp}) -> store now holds {len(store['snapshots'])} snapshot(s)")
    print(f"  rotations={pat['rotations']}  routes={len(pat['weekly_seats_by_destination'])}  "
          f"fleet={len(pat['weekly_movements_by_aircraft'])} types")


def show():
    store = load_store()
    print(f"pattern store: {len(store['snapshots'])} snapshot(s)")
    for s in store['snapshots']:
        p = s['patterns']
        print(f"  {s['ingested_at']}  {s['file']}  ({s['sha16']})  "
              f"rotations={p['rotations']}")
    cur = store['current'].get('schedule')
    if cur:
        p = cur['patterns']
        top = list(p['weekly_seats_by_destination'].items())[:3]
        print("current schedule pattern:", ', '.join(f"{k} {v:,} seats/wk" for k, v in top))


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == 'ingest':
        ingest(sys.argv[2])
    elif len(sys.argv) >= 2 and sys.argv[1] == 'show':
        show()
    else:
        ingest(DEFAULT_DROP)      # seed with the S26 schedule
        show()
